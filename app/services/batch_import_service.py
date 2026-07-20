"""
Excel (.xlsx) batch upload / bulk import support:
  - Column definitions for the two templates (single-submission line items,
    and multi-submission bulk import).
  - Template generation (openpyxl workbook -> bytes).
  - Upload parsing (openpyxl workbook -> raw row dicts, matched by header label
    so re-ordered columns still work).
  - Row coercion into the plain-dict shape LineItemIn/SubmissionIn expect.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Column definitions: (field_name, header_label, required)
# ---------------------------------------------------------------------------

LINE_ITEM_COLUMNS: list[tuple[str, str, bool]] = [
    ("vendor_name", "Vendor Name", True),
    ("invoice_no", "Invoice No", True),
    ("po_number", "PO Number", False),
    ("description", "Description", True),
    ("items_products", "Items/Products", False),
    ("category", "Category", True),
    ("account_code", "Account Code", True),
    ("billing_period_start", "Billing Period Start (YYYY-MM-DD)", True),
    ("billing_period_end", "Billing Period End (YYYY-MM-DD)", True),
    ("payment_tracking_code", "Payment Tracking Code", False),
    ("frequency", "Frequency (one_off/monthly/quarterly/annual)", True),
    ("status_remarks", "Status Remarks", False),
    ("currency", "Currency (USD/NGN/EUR/GBP/INR)", True),
    ("original_amount", "Original Amount", True),
    ("is_arrear", "Is Arrear (TRUE/FALSE)", False),
    ("arrear_type", "Arrear Type", False),
]

LINE_ITEM_EXAMPLE: dict[str, Any] = {
    "vendor_name": "Acme Supplies Ltd",
    "invoice_no": "INV-2026-001",
    "po_number": "PO-4521",
    "description": "Monthly maintenance contract",
    "items_products": "Spare filters",
    "category": "Maintenance Cost",
    "account_code": "6100-01",
    "billing_period_start": "2026-07-01",
    "billing_period_end": "2026-07-31",
    "payment_tracking_code": "",
    "frequency": "monthly",
    "status_remarks": "",
    "currency": "USD",
    "original_amount": 12500.00,
    "is_arrear": "FALSE",
    "arrear_type": "",
}

BULK_IMPORT_COLUMNS: list[tuple[str, str, bool]] = [
    ("batch_ref", "Batch Ref (same value groups rows into one submission)", True),
    ("department", "Department", True),
    ("month", "Month (1-12)", True),
    ("year", "Year", True),
    ("cost_type", "Cost Type (opex/capex)", True),
    ("supporting_justification", "Supporting Justification", True),
] + LINE_ITEM_COLUMNS

BULK_IMPORT_EXAMPLE: dict[str, Any] = {
    "batch_ref": "BATCH-1",
    "department": "DPRP-Information Technology",
    "month": 7,
    "year": 2026,
    "cost_type": "opex",
    "supporting_justification": "Routine monthly operating costs",
    **LINE_ITEM_EXAMPLE,
}


# ---------------------------------------------------------------------------
# Template generation
# ---------------------------------------------------------------------------

def build_template_workbook(
    columns: list[tuple[str, str, bool]],
    example_row: dict[str, Any] | None = None,
) -> bytes:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    ws.title = "Data"

    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, (_, label, required) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label + (" *" if required else ""))
        cell.font = header_font
        cell.fill = openpyxl.styles.PatternFill("solid", fgColor="003087")
        ws.column_dimensions[cell.column_letter].width = max(18, min(40, len(label) // 1.3))

    if example_row:
        for col_idx, (field, _, _) in enumerate(columns, start=1):
            ws.cell(row=2, column=col_idx, value=example_row.get(field, ""))

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Upload parsing
# ---------------------------------------------------------------------------

def parse_uploaded_workbook(
    file_bytes: bytes,
    columns: list[tuple[str, str, bool]],
) -> tuple[list[dict[str, Any]], list[str]]:
    """
    Returns (raw_rows, errors).
    raw_rows: one dict per data row, keyed by field_name, values as read from
              the cell (still un-coerced — dates may be datetime, numbers may
              be int/float/str depending on how Excel stored them).
    errors:   top-level parse errors (bad file, missing columns, empty file).
              If non-empty, raw_rows is always [].
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        return [], [f"Could not read the uploaded file — please ensure it is a valid .xlsx file. ({exc})"]

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], ["The uploaded file is empty."]

    normalized = [
        str(h).strip().lower().rstrip(" *") if h is not None else ""
        for h in header_row
    ]

    header_map: dict[str, int] = {}
    errors: list[str] = []
    for field, label, _required in columns:
        target = label.strip().lower()
        try:
            idx = normalized.index(target)
        except ValueError:
            errors.append(f'Missing expected column: "{label}". Please use the provided template.')
            continue
        header_map[field] = idx

    if errors:
        return [], errors

    raw_rows: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None or all(
            c is None or (isinstance(c, str) and not c.strip()) for c in row
        ):
            continue
        raw: dict[str, Any] = {}
        for field, _, _ in columns:
            idx = header_map[field]
            raw[field] = row[idx] if idx < len(row) else None
        raw_rows.append(raw)

    return raw_rows, []


# ---------------------------------------------------------------------------
# Row coercion — raw cell values -> plain dict compatible with LineItemIn
# ---------------------------------------------------------------------------

def _s(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _parse_date(v: Any, field_label: str) -> date:
    if v is None or v == "":
        raise ValueError(f"{field_label} is required.")
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    try:
        return date.fromisoformat(s[:10])
    except ValueError:
        raise ValueError(f"{field_label}: could not parse date '{v}'. Use YYYY-MM-DD.")


def _parse_bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    return str(v).strip().lower() in ("true", "yes", "y", "1")


def _parse_amount(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, str):
        return v.replace(",", "").strip()
    return v


def coerce_line_item_row(raw: dict[str, Any]) -> dict[str, Any]:
    """
    Raises ValueError with a human-readable message on the first hard
    problem (missing required date). Everything else is left for the
    LineItemIn pydantic validators to catch and report.
    """
    return {
        "vendor_name": _s(raw.get("vendor_name")) or "",
        "invoice_no": _s(raw.get("invoice_no")) or "",
        "po_number": _s(raw.get("po_number")),
        "description": _s(raw.get("description")) or "",
        "items_products": _s(raw.get("items_products")),
        "category": _s(raw.get("category")) or "",
        "account_code": _s(raw.get("account_code")) or "",
        "billing_period_start": _parse_date(raw.get("billing_period_start"), "Billing Period Start"),
        "billing_period_end": _parse_date(raw.get("billing_period_end"), "Billing Period End"),
        "payment_tracking_code": _s(raw.get("payment_tracking_code")),
        "frequency": (_s(raw.get("frequency")) or "").lower().replace(" ", "_"),
        "status_remarks": _s(raw.get("status_remarks")),
        "currency": (_s(raw.get("currency")) or "").upper(),
        "original_amount": _parse_amount(raw.get("original_amount")),
        "is_arrear": _parse_bool(raw.get("is_arrear")),
        "arrear_type": _s(raw.get("arrear_type")),
    }
